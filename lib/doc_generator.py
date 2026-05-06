#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
doc_generator.py — Word 文档生成模块
从 gensjt.py 提取，负责纯净题目卷和完整答案卷的生成。
"""

import os
import random
from collections import defaultdict

import docx
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn


def _normalize_items(items):
    """
    将分批JSON格式或分组格式统一转换为 [(dim_name, [items...]), ...] 格式。
    分批JSON格式: [{"dimension": "...", "stem": "...", "options": {...}}, ...]
    分组格式:     [(dim_name, [item, ...]), ...]
    """
    if not items:
        return []
    first = items[0]
    if isinstance(first, tuple):
        return items
    groups = defaultdict(list)
    for item in items:
        dim = item.get("dimension", "未知维度")
        groups[dim].append(item)
    return [(dim, list(items_list)) for dim, items_list in groups.items()]


def _add_run(para, text, font_name="SimSun", font_size=Pt(12), bold=False):
    """向段落添加格式化文本的便捷方法"""
    run = para.add_run(text)
    run.font.name = font_name
    run.font.size = font_size
    run.font.bold = bold
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    return run


def create_test_doc(items: list, industry: str, position: str, output_path: str):
    """
    生成纯净题目卷（选项打乱）。
    打乱后按位置重新标记 A/B/C/D，映射存入 item['test_labels']。
    """
    doc = docx.Document()

    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = section.right_margin = section.top_margin = section.bottom_margin = Cm(2.54)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_run(title, "情境判断测验", font_size=Pt(18), bold=True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_run(subtitle, f"【{industry}】【{position}】", font_size=Pt(12))

    doc.add_paragraph()

    for dim_idx, (dim_name, dim_items) in enumerate(items, 1):
        dim_title = doc.add_paragraph()
        _add_run(dim_title, f"维度{dim_idx}：{dim_name}", font_size=Pt(12), bold=True)

        for q_idx, item in enumerate(dim_items, 1):
            q_para = doc.add_paragraph()
            _add_run(q_para, f"{q_idx}. {item['stem']}")

            raw_options = [
                ("A", item['options']['A']),
                ("B", item['options']['B']),
                ("C", item['options']['C']),
                ("D", item['options']['D']),
            ]
            random.seed(q_idx * 100 + dim_idx * 17)
            random.shuffle(raw_options)

            pos_labels = ["A", "B", "C", "D"]
            test_labels = {}
            for pos_idx, (orig_key, _) in enumerate(raw_options):
                test_labels[pos_labels[pos_idx]] = orig_key

            item["test_labels"] = test_labels

            for pos_idx, (orig_key, opt_data) in enumerate(raw_options):
                opt_para = doc.add_paragraph()
                opt_para.paragraph_format.left_indent = Cm(0.5)
                _add_run(opt_para, f"{pos_labels[pos_idx]} {opt_data['text']}", font_size=Pt(11))

            doc.add_paragraph()

    os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)
    doc.save(output_path)


def create_answer_doc(items: list, industry: str, position: str, output_path: str):
    """
    生成完整答案卷。
    选项标签与测试卷保持一致（使用 item['test_labels'] 映射），
    按分值 0→1→2 升序排列（2-1-1-0 分制）。
    """
    doc = docx.Document()

    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = section.right_margin = section.top_margin = section.bottom_margin = Cm(2.54)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_run(title, "情境判断测验（答案卷）", font_size=Pt(18), bold=True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_run(subtitle, f"【{industry}】【{position}】", font_size=Pt(12))

    doc.add_paragraph()

    for dim_idx, (dim_name, dim_items) in enumerate(items, 1):
        dim_title = doc.add_paragraph()
        _add_run(dim_title, f"维度{dim_idx}：{dim_name}", font_size=Pt(12), bold=True)

        for q_idx, item in enumerate(dim_items, 1):
            q_para = doc.add_paragraph()
            _add_run(q_para, f"{q_idx}. {item['stem']}")

            test_labels = item.get("test_labels")
            if test_labels:
                orig_to_pos = {v: k for k, v in test_labels.items()}
            else:
                orig_to_pos = {"A": "A", "B": "B", "C": "C", "D": "D"}

            sorted_opts = sorted(
                item['options'].items(),
                key=lambda kv: kv[1]['score']
            )

            for orig_key, opt_data in sorted_opts:
                pos_label = orig_to_pos[orig_key]
                score = opt_data['score']
                opt_para = doc.add_paragraph()
                opt_para.paragraph_format.left_indent = Cm(0.5)
                _add_run(opt_para, f"{pos_label}（{score}分）{opt_data['text']}", font_size=Pt(11))

            doc.add_paragraph()

            explain_title = doc.add_paragraph()
            _add_run(explain_title, "【赋分说明】", font_size=Pt(10.5), bold=True)

            for orig_key, opt_data in sorted_opts:
                pos_label = orig_to_pos[orig_key]
                score = opt_data['score']
                p = doc.add_paragraph()
                _add_run(p, f"{pos_label}={score}分：{opt_data['reason']}", font_size=Pt(10.5))

            doc.add_paragraph()

    os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)
    doc.save(output_path)
