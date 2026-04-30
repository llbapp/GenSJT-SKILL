#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
excel_generator.py — Excel 汇总表生成模块
从 gensjt.py 提取，负责生成包含3个Sheet的 Excel 汇总表。
"""

import os
from collections import defaultdict

import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill


def create_excel_summary(items: list, dimension_defs: dict, industry: str,
                         position: str, output_path: str,
                         knowledge_refs: dict = None,
                         search_context_fn=None):
    """
    生成Excel汇总表（3个Sheet：题目明细、维度描述、知识库引用记录）。

    参数:
        items: 题目列表（分组或单条格式均可）
        dimension_defs: 维度定义字典
        industry: 行业名称
        position: 岗位名称
        output_path: 输出文件路径
        knowledge_refs: 知识库引用信息
        search_context_fn: 行业岗位检索函数（可选，用于Sheet3）。
                          签名: (industry, position, top_n) -> dict
    """
    wb = openpyxl.Workbook()

    thin = Side(style="thin")
    thin_border = Border(left=thin, right=thin, top=thin, bottom=thin)
    hdr_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    hdr_font = Font(name="微软雅黑", bold=True, size=10)

    _build_sheet1(wb, items, industry, position, thin_border, hdr_fill, hdr_font)
    _build_sheet2(wb, items, dimension_defs, thin_border, hdr_fill, hdr_font)
    _build_sheet3(wb, items, industry, position, knowledge_refs,
                  thin_border, hdr_fill, hdr_font, search_context_fn)

    os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)
    wb.save(output_path)


# ── Sheet1: 题目明细 ──
def _build_sheet1(wb, items, industry, position, thin_border, hdr_fill, hdr_font):
    ws = wb.active
    ws.title = "题目明细"

    headers = ["序号", "行业", "岗位", "维度", "题干",
               "选项A", "选项B", "选项C", "选项D",
               "A分值", "B分值", "C分值", "D分值",
               "A说明", "B说明", "C说明", "D说明",
               "P值(难度)", "P难度等级", "D值(区分度)", "D区分度等级", "P/D估算理由"]

    ws.append(headers)
    for col in range(1, len(headers) + 1):
        c = ws.cell(row=1, column=col)
        c.font = hdr_font
        c.fill = hdr_fill
        c.border = thin_border
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    row_num = 2
    dim_list = _ensure_grouped(items)
    for dim_idx, (dim_name, dim_items) in enumerate(dim_list, 1):
        for q_idx, item in enumerate(dim_items, 1):
            row_data = [
                f"{dim_idx}.{q_idx}", industry, position, dim_name,
                item['stem'],
                item['options']['A']['text'], item['options']['B']['text'],
                item['options']['C']['text'], item['options']['D']['text'],
                0, 1, 2, 3,
                item['options']['A']['reason'], item['options']['B']['reason'],
                item['options']['C']['reason'], item['options']['D']['reason'],
                item.get('p_value', ''), item.get('p_level', ''),
                item.get('d_value', ''), item.get('d_level', ''),
                item.get('p_d_reason', '')
            ]
            ws.append(row_data)
            for col in range(1, len(headers) + 1):
                c = ws.cell(row=row_num, column=col)
                c.border = thin_border
                c.alignment = Alignment(vertical="top", wrap_text=True)
                c.font = Font(
                    name="微软雅黑" if col <= 4 or col == 5 or col >= 12 else "Times New Roman",
                    size=10
                )
            row_num += 1

    col_widths = [8, 10, 12, 14, 55, 38, 38, 38, 38, 7, 7, 7, 7, 45, 45, 45, 45, 8, 10, 8, 10, 45]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w


# ── Sheet2: 维度描述 ──
def _build_sheet2(wb, items, dimension_defs, thin_border, hdr_fill, hdr_font):
    ws = wb.create_sheet("维度描述")

    headers = ["维度", "定义", "高分表现", "低分表现",
               "关键行为指标1", "关键行为指标2", "关键行为指标3",
               "关键行为指标4", "关键行为指标5",
               "优秀(5)", "良好(4)", "中等(3)", "欠佳(2)", "不足(1)"]

    ws.append(headers)
    for col in range(1, len(headers) + 1):
        c = ws.cell(row=1, column=col)
        c.font = hdr_font
        c.fill = hdr_fill
        c.border = thin_border
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    dim_list = _ensure_grouped(items)
    for dim_idx, (dim_name, _) in enumerate(dim_list, 1):
        dim_info = dimension_defs.get(dim_name, {})
        kbi = dim_info.get("关键行为指标", [])
        if isinstance(kbi, str):
            kbi = [kbi] + [""] * 4
        kbi = (kbi + [""] * 5)[:5]

        row_data = [
            dim_name,
            dim_info.get("定义", ""),
            dim_info.get("高分表现", ""),
            dim_info.get("低分表现", ""),
        ] + kbi + [
            dim_info.get("优秀", ""),
            dim_info.get("良好", ""),
            dim_info.get("中等", ""),
            dim_info.get("欠佳", ""),
            dim_info.get("不足", "")
        ]
        ws.append(row_data)
        for col in range(1, len(headers) + 1):
            c = ws.cell(row=dim_idx + 1, column=col)
            c.border = thin_border
            c.alignment = Alignment(vertical="top", wrap_text=True)
            c.font = Font(name="微软雅黑", size=10)

    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 45
    ws.column_dimensions["C"].width = 35
    ws.column_dimensions["D"].width = 35
    for col_letter in ["E", "F", "G", "H", "I"]:
        ws.column_dimensions[col_letter].width = 32
    for col_letter in ["J", "K", "L", "M", "N"]:
        ws.column_dimensions[col_letter].width = 38


# ── Sheet3: 知识库引用记录 ──
def _build_sheet3(wb, items, industry, position, knowledge_refs,
                  thin_border, hdr_fill, hdr_font, search_context_fn):
    ws = wb.create_sheet("知识库引用记录")

    headers = ["序号", "维度",
               "胜任特征辞典", "辞典引用说明",
               "母题模板ID",
               "胜任-情境对应", "情境对应说明",
               "例题库参考", "例题参考说明",
               "行业岗位参考", "岗位匹配说明"]

    ws.append(headers)
    for col in range(1, len(headers) + 1):
        c = ws.cell(row=1, column=col)
        c.font = hdr_font
        c.fill = hdr_fill
        c.border = thin_border
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    dim_list = _ensure_grouped(items)

    # 行业岗位匹配信息（仅计算一次）
    job_ctx_info = ""
    job_ctx_detail = ""
    if industry and position and search_context_fn:
        try:
            ctx = search_context_fn(industry, position, top_n=3)
            if ctx.get("results"):
                top_match = ctx["results"][0]
                meta = top_match.get("_meta", {})
                matched_ind = top_match.get("industry", "")
                matched_pos = top_match.get("position_archetype", "")
                matched_cat = top_match.get("position_category", "")
                combined_score = meta.get("match_score", 0)
                ind_score = meta.get("industry_score", 0)
                pos_score = meta.get("position_score", 0)
                job_ctx_info = "✓ 已匹配"
                job_ctx_detail = (
                    f"用户输入：{industry} / {position}\n"
                    f"最佳匹配：{matched_ind} / {matched_cat} - {matched_pos}\n"
                    f"匹配分：综合 {combined_score}（行业 {ind_score} + 岗位 {pos_score}）"
                )
                if len(ctx["results"]) > 1:
                    runners = []
                    for r in ctx["results"][1:]:
                        rmeta = r.get("_meta", {})
                        runners.append(
                            f"{r.get('industry', '')} / {r.get('position_archetype', '')}"
                            f"（{rmeta.get('match_score', 0)}）"
                        )
                    job_ctx_detail += f"\n备选：{'；'.join(runners)}"
            else:
                job_ctx_info = "✗ 未匹配"
                job_ctx_detail = (f"用户输入：{industry} / {position}，"
                                  f"未在行业岗位知识库中找到匹配记录")
        except Exception:
            job_ctx_info = "✗ 检索失败"
            job_ctx_detail = ""

    row_num = 2
    for dim_idx, (dim_name, dim_items) in enumerate(dim_list, 1):
        ref_info = (knowledge_refs or {}).get(dim_name, {})
        dict_found = ref_info.get("dict_found", False)
        dict_summary = ref_info.get("dict_summary", "")
        template_ids = ref_info.get("template_ids", [])
        competence_found = ref_info.get("competence_found", False)
        competence_summary = ref_info.get("competence_summary", "")
        examples_found = ref_info.get("examples_found", False)
        examples_summary = ref_info.get("examples_summary", "")

        for q_idx, item in enumerate(dim_items, 1):
            row_data = [
                f"{dim_idx}.{q_idx}",
                dim_name,
                "✓ 已匹配" if dict_found else "✗ 未匹配",
                dict_summary,
                ", ".join(template_ids) if template_ids else "无",
                "✓ 已匹配" if competence_found else "✗ 未匹配",
                competence_summary,
                "✓ 已匹配" if examples_found else "✗ 未匹配",
                examples_summary,
                job_ctx_info if job_ctx_info else "未提供行业岗位",
                job_ctx_detail if job_ctx_info else "",
            ]
            ws.append(row_data)
            for col in range(1, len(headers) + 1):
                c = ws.cell(row=row_num, column=col)
                c.border = thin_border
                c.alignment = Alignment(vertical="top", wrap_text=True)
                c.font = Font(name="微软雅黑", size=10)
            row_num += 1

    col_widths = [8, 14, 12, 45, 20, 14, 45, 12, 45, 14, 55]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w


# ── 辅助函数 ──
def _ensure_grouped(items):
    """将任意格式（分组/单条）统一为 [(dim_name, [items...]), ...] 分组格式"""
    if items and isinstance(items[0], tuple):
        return items
    dim_groups = defaultdict(list)
    for item in items:
        dim = item.get("dimension", "未知维度")
        dim_groups[dim].append(item)
    return list(dim_groups.items())
